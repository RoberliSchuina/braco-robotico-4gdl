# -*- coding: utf-8 -*-
"""gerar_lista_parafusos.py — gera o documento à parte relatorio/Lista_Parafusos_Fixadores.docx (lista de compras)."""
import os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parafusos import FIXADORES, NOTAS, custo_total

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "relatorio", "Lista_Parafusos_Fixadores.docx")

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
sec.left_margin = sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(1.5)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

h = doc.add_heading("Lista de parafusos e fixadores — Braço Robótico 4 GDL", 1)
h.runs[0].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
p = doc.add_paragraph("Sistemas Embarcados — FAESA — 2026/2. Lista de compras à parte da BOM eletrônica, dimensionada pelos "
                      "furos das peças impressas (cad/gerar_pecas.py) e pelos empilhamentos de montagem.")
p.runs[0].italic = True

cab = ["#", "Especificação (modelo)", "Necess.", "Comprar", "Onde é usado", "Furo na peça / empilhamento", "Material sugerido", "R$ est."]
larg = [0.7, 5.2, 1.3, 1.8, 6.0, 6.0, 3.0, 1.2]
t = doc.add_table(rows=1, cols=len(cab)); t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, x in zip(t.rows[0].cells, cab):
    c.text = ""; r = c.paragraphs[0].add_run(x); r.bold = True; r.font.size = Pt(8.5)
for item in FIXADORES:
    cells = t.add_row().cells
    for c, v in zip(cells, item):
        c.text = ""; r = c.paragraphs[0].add_run(str(v)); r.font.size = Pt(8)
for row in t.rows:
    for c, w in zip(row.cells, larg): c.width = Cm(w)

p = doc.add_paragraph(); r = p.add_run(f"Custo total estimado dos fixadores: R$ {custo_total():.0f} (preços de referência 2026; confirmar na compra)."); r.bold = True

doc.add_heading("Resumo por tipo (para o balcão da loja)", 2)
resumo = [
    ("Autoatarraxante PA 2,0 × 8 mm", "20 un."), ("Micro autoatarraxante PA 1,7 × 6 mm", "kit 50 un."),
    ("M3 × 25 mm inox", "2 un."), ("Porca M3 autotravante (nylock)", "2 un."), ("Arruela M3 (2 nylon + 2 inox)", "4 un."),
    ("M3 × 12 mm", "10 un."), ("Porca M3 sextavada", "10 un."), ("Arruela M3 ampla Ø9", "10 un."),
    ("M3 × 16 mm + porca (ou p/ madeira 3,0 × 16)", "4 un."), ("M3 × 6 mm + espaçador nylon M3 × 6", "4 un."),
]
t2 = doc.add_table(rows=1, cols=2); t2.style = "Light Grid Accent 1"; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, x in zip(t2.rows[0].cells, ["Tipo", "Quantidade"]):
    c.text = ""; r = c.paragraphs[0].add_run(x); r.bold = True; r.font.size = Pt(9)
for a, b in resumo:
    cells = t2.add_row().cells; cells[0].text = a; cells[1].text = b
    for c in cells: c.paragraphs[0].runs[0].font.size = Pt(9)
for row in t2.rows:
    row.cells[0].width = Cm(9); row.cells[1].width = Cm(4)

doc.add_heading("Notas", 2)
for n in NOTAS:
    doc.add_paragraph(n, style="List Bullet")

fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Braço Robótico 4 GDL — Lista de parafusos e fixadores — gerada por relatorio/gerar_lista_parafusos.py").font.size = Pt(8)
try:
    doc.save(SAIDA); print("OK", SAIDA)
except PermissionError:
    alt = SAIDA.replace(".docx", "_novo.docx"); doc.save(alt); print("AVISO: arquivo aberto no Word; salvo como", alt)
